"""Build the Thai document as DOCX (python-docx) or PDF (HTML rendered by WeasyPrint)."""
import io
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from jinja2 import Environment, PackageLoader, select_autoescape

from . import config, tables

BULLET = re.compile(r"^[-•·*‣◦▪–]$")
_env = Environment(loader=PackageLoader("app", "templates"), autoescape=select_autoescape())


def block_text(block):
    """Translated text when there is one; otherwise the source (verbatim or not yet translated)."""
    if block["status"] == "done" and block["translation"]:
        return block["translation"]
    return block["source"]


def _is_bullet(marker):
    return not marker or bool(BULLET.match(marker.strip()))


# --- DOCX ---------------------------------------------------------------------------------

def _set_style_font(style, font, size_pt):
    """Thai is a 'complex script' in Word: it is drawn with the w:cs font at the w:szCs size,
    so both have to be set, and theme fonts removed or they override the explicit names."""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)
    style.font.size = Pt(size_pt)
    # Word rejects rPr children out of schema order, so insert before the elements that must follow.
    szcs = _get_or_insert(rpr, "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
                          "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                          "w:specVanish", "w:oMath")
    szcs.set(qn("w:val"), str(round(size_pt * 2)))
    lang = _get_or_insert(rpr, "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
    lang.set(qn("w:bidi"), "th-TH")


def _get_or_insert(parent, tag, *successors):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.insert_element_before(el, *successors)
    return el


def build_docx(blocks, title):
    doc = Document()
    font, size = config.DOCX_FONT, config.DOCX_FONT_SIZE
    _set_style_font(doc.styles["Normal"], font, size)
    _set_style_font(doc.styles["Caption"], font, size - 1)
    for level, extra in ((1, 6), (2, 4), (3, 2), (4, 1)):
        _set_style_font(doc.styles[f"Heading {level}"], font, size + extra)
    doc.core_properties.title = title

    for b in blocks:
        text, kind = block_text(b), b["kind"]
        if kind == "heading":
            doc.add_heading(text, level=min(max(b["level"] or 1, 1), 4))
        elif kind == "list_item":
            if _is_bullet(b["marker"]):
                doc.add_paragraph(text, style="List Bullet")
            else:
                # Write the original number/letter; Word's auto-numbering would renumber across lists.
                p = doc.add_paragraph(f"{b['marker']} {text}")
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.first_line_indent = Pt(-18)
        elif kind == "table":
            rows = tables.parse(text)
            if not rows:
                doc.add_paragraph(text)
                continue
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for col, cell_text in enumerate(row):
                    cell = table.cell(r, col)
                    cell.text = cell_text
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif kind in ("caption", "footnote"):
            doc.add_paragraph(text, style="Caption")
        elif kind in ("code", "formula"):
            p = doc.add_paragraph()
            for i, line in enumerate(text.splitlines()):
                if i:
                    p.add_run().add_break(WD_BREAK.LINE)
                p.add_run(line).font.name = "Consolas"
        else:
            doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF ----------------------------------------------------------------------------------

def _html_items(blocks):
    for b in blocks:
        item = {"kind": b["kind"], "text": block_text(b), "level": min(max(b["level"] or 1, 1), 4),
                "marker": None if _is_bullet(b["marker"]) else b["marker"]}
        if b["kind"] == "table":
            item["rows"] = tables.parse(item["text"])
        yield item


def build_html(blocks, title):
    return _env.get_template("export.html").render(title=title, items=list(_html_items(blocks)), font=config.PDF_FONT)


def build_pdf(blocks, title):
    from weasyprint import HTML  # imported lazily: only needed for PDF export

    return HTML(string=build_html(blocks, title)).write_pdf()
