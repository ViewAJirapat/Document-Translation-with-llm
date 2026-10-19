import io

import pypdfium2
from docx import Document
from docx.oxml.ns import qn

from app import config, export

BLOCKS = [
    {"kind": "heading", "level": 1, "marker": None, "status": "done", "source": "Introduction", "translation": "บทนำ"},
    {"kind": "paragraph", "level": None, "marker": None, "status": "done",
     "source": "Floods are common.", "translation": "น้ำท่วมเกิดขึ้นบ่อยในกรุงเทพมหานครในช่วงฤดูฝน " * 8},
    {"kind": "list_item", "level": None, "marker": "-", "status": "done", "source": "One", "translation": "หนึ่ง"},
    {"kind": "list_item", "level": None, "marker": "2.", "status": "done", "source": "Two", "translation": "สอง"},
    {"kind": "table", "level": None, "marker": None, "status": "done", "source": "| A | B |\n|---|---|\n| x | 1 |",
     "translation": "| ก | ข |\n|---|---|\n| เอ็กซ์ | 1 |"},
    {"kind": "paragraph", "level": None, "marker": None, "status": "pending", "source": "Not yet translated.", "translation": None},
    {"kind": "code", "level": None, "marker": None, "status": "skip", "source": "print('hi')\nexit()", "translation": None},
]


def test_docx_contains_thai_text_table_and_fallbacks():
    doc = Document(io.BytesIO(export.build_docx(BLOCKS, "Test")))
    texts = [p.text for p in doc.paragraphs]
    assert "บทนำ" in texts
    assert "2. สอง" in texts  # numbered markers are written literally
    assert "Not yet translated." in texts  # untranslated blocks fall back to the English
    assert doc.tables[0].cell(1, 0).text == "เอ็กซ์"


def test_docx_sets_complex_script_font_in_schema_order():
    doc = Document(io.BytesIO(export.build_docx(BLOCKS, "Test")))
    for name in ("Normal", "Heading 1", "Caption"):
        rpr = doc.styles[name].element.rPr
        rfonts = rpr.find(qn("w:rFonts"))
        assert rfonts.get(qn("w:cs")) == config.DOCX_FONT
        assert rfonts.get(qn("w:cstheme")) is None
        assert rpr.find(qn("w:lang")).get(qn("w:bidi")) == "th-TH"
        tags = [child.tag for child in rpr]
        # Word requires rFonts first and sz < szCs < lang (CT_RPr schema order).
        assert tags[0] == qn("w:rFonts")
        assert tags.index(qn("w:sz")) < tags.index(qn("w:szCs")) < tags.index(qn("w:lang"))


def test_pdf_renders_thai_text():
    data = export.build_pdf(BLOCKS, "Test")
    assert data.startswith(b"%PDF")
    pdf = pypdfium2.PdfDocument(data)
    text = "".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
    # WeasyPrint draws Thai correctly, but its text layer maps glyphs shared between clusters
    # (e.g. the า in ำ) to one string, so copied Thai text is imperfect: check characters, not words.
    assert sum("฀" <= ch <= "๿" for ch in text) > 100
    assert "Not yet translated." in text


def test_html_marks_document_as_thai_and_uses_configured_font():
    html = export.build_html(BLOCKS, "Test")
    assert '<html lang="th">' in html
    assert config.PDF_FONT in html
